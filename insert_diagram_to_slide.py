"""
insert_diagram_to_slide.py
Cleans Slide 12 (removing the old text cards that were covering the slide)
and inserts the high-resolution use_case_diagram.png right into Slide 12.
"""

from pptx import Presentation
from pptx.util import Inches
import os

IMG_PATH = r"D:\project\FSD\crop-disease-prediction\use_case_diagram.png"
PPT_FILES = [
    r"D:\project\FSD\crop-disease-prediction\FSD PROJECT PPT FINAL_Updated.pptx",
    r"D:\project\FSD\crop-disease-prediction\FSD PROJECT PPT FINAL_WithUseCase.pptx",
]

for ppt_path in PPT_FILES:
    if not os.path.exists(ppt_path):
        continue
    
    prs = Presentation(ppt_path)
    slide = prs.slides[11]  # Slide 12 is 0-indexed at 11
    
    # Identify shapes to keep: Background rectangle (shape 0) and Title Bar shapes
    # Any shape that was previously added as a card, text block, or previous picture will be removed.
    shapes_to_remove = []
    for idx, shape in enumerate(slide.shapes):
        # We keep the background (usually idx 0) and the title header (idx 1 and 2)
        # Let's check text: if it says "Use Case Diagram", keep it and its backing bar
        if hasattr(shape, "text") and "Use Case Diagram" in shape.text:
            continue
        # Also keep title bar rectangle (top around 0.54 in, height around 0.72 in)
        if shape.top < Inches(1.30):
            continue
        shapes_to_remove.append(shape)
    
    for shape in shapes_to_remove:
        sp = shape._element
        sp.getparent().remove(sp)
        
    print(f"[{os.path.basename(ppt_path)}] Removed {len(shapes_to_remove)} old overlapping card shapes from Slide 12.")
    
    # Now insert the diagram picture, perfectly sized and centered in the slide content area
    # Slide width = 10.0 in, height = 5.625 in
    # Title bar ends at ~1.26 in
    left = Inches(0.50)
    top = Inches(1.35)
    width = Inches(9.00)
    height = Inches(4.09)
    
    pic = slide.shapes.add_picture(IMG_PATH, left, top, width, height)
    print(f"[{os.path.basename(ppt_path)}] Added use_case_diagram.png: left={left.inches}in, top={top.inches}in, width={width.inches}in, height={height.inches}in")
    
    prs.save(ppt_path)
    print(f"[{os.path.basename(ppt_path)}] Successfully saved!\n")
